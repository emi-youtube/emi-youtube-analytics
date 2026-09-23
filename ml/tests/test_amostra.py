"""Testes da amostragem humana e das planilhas dos avaliadores.

Nada toca o banco: a lógica que importa (Cochran, estratificação, embaralhamento,
montagem da planilha) é pura e testável com dados sintéticos.
"""

import asyncio
import inspect
import random

import pytest
from openpyxl import load_workbook

from ml.amostra.gerar_planilhas_avaliadores import AVALIADORES, CABECALHO, montar_planilha
from ml.amostra.sortear_amostra_humana import (
    alocar_por_estrato,
    carregar_sorteaveis,
    contar_rotulados,
    sortear,
    tamanho_cochran,
)
from ml.config import CLASSES, SEMENTE

# ------------------------------------------------------------------- Cochran


def test_cochran_para_o_corpus_da_sprint1():
    """N = 2.534 (corpus da execucao 4) deve dar n = 334."""
    assert tamanho_cochran(2534) == 334


def test_cochran_converge_para_385_em_populacao_grande():
    """Sem correcao finita, n0 = 1,96^2 * 0,25 / 0,05^2 = 384,16 -> 385."""
    assert tamanho_cochran(10**9) == 385


def test_cochran_nunca_passa_da_populacao():
    for n in (10, 50, 100, 334, 2534):
        assert tamanho_cochran(n) <= n


def test_cochran_cresce_com_a_populacao():
    tamanhos = [tamanho_cochran(n) for n in (100, 500, 1000, 2534, 10000)]
    assert tamanhos == sorted(tamanhos)


# -------------------------------------------------------------- estratificação


def test_alocacao_divide_em_tercos_quando_ha_folga():
    cotas = alocar_por_estrato(334, {"positivo": 1000, "negativo": 1000, "neutro": 1000})
    assert sum(cotas.values()) == 334
    # 334 nao divide por 3: a sobra fica com os primeiros da rodada
    assert sorted(cotas.values()) == [111, 111, 112]


def test_estrato_pequeno_leva_todos_e_a_sobra_e_redistribuida():
    """Estrato menor que a cota leva todos, e a sobra vai para os outros: n nao encolhe."""
    cotas = alocar_por_estrato(334, {"positivo": 1500, "negativo": 40, "neutro": 900})
    assert cotas["negativo"] == 40
    assert sum(cotas.values()) == 334


def test_alocacao_nunca_pede_mais_do_que_existe():
    disponivel = {"positivo": 20, "negativo": 5, "neutro": 7}
    cotas = alocar_por_estrato(334, disponivel)
    for classe, cota in cotas.items():
        assert cota <= disponivel[classe]
    assert sum(cotas.values()) == 32


def test_classe_ausente_nao_entra_na_alocacao():
    cotas = alocar_por_estrato(300, {"positivo": 500, "neutro": 500})
    assert "negativo" not in cotas
    assert sum(cotas.values()) == 300


def test_alocacao_e_deterministica():
    disponivel = {"positivo": 1200, "negativo": 380, "neutro": 954}
    assert alocar_por_estrato(334, disponivel) == alocar_por_estrato(334, disponivel)


# ---------------------------------------------------------------- sorteio


def test_sorteio_com_semente_fixa_e_reproduzivel():
    ids = list(range(1, 1001))
    a = random.Random(SEMENTE).sample(ids, 112)
    b = random.Random(SEMENTE).sample(ids, 112)
    assert a == b


def test_cada_avaliador_recebe_ordem_diferente():
    """Ordem igual nos tres faria os avaliadores cansarem nos mesmos comentarios."""
    linhas = [(i, f"comentario {i}") for i in range(1, 201)]
    ordens = []
    for indice in range(len(AVALIADORES)):
        embaralhado = linhas[:]
        random.Random(SEMENTE + indice).shuffle(embaralhado)
        ordens.append([id_ for id_, _ in embaralhado])

    for i in range(len(ordens)):
        for j in range(i + 1, len(ordens)):
            assert ordens[i] != ordens[j], f"avaliadores {i} e {j} receberam a mesma ordem"
    # mesmo conteudo, so a ordem muda
    for ordem in ordens:
        assert sorted(ordem) == [id_ for id_, _ in linhas]


# ---------------------------------------------------------------- planilha


@pytest.fixture
def planilha(tmp_path):
    linhas = [(10, "amei 😂"), (11, "que lixo"), (12, "quando lanca?")]
    caminho = tmp_path / "a.xlsx"
    montar_planilha(linhas, "avaliador_1").save(caminho)
    return load_workbook(caminho)


def test_planilha_tem_as_colunas_combinadas(planilha):
    aba = planilha["rotulagem"]
    cabecalho = tuple(c.value for c in aba[1])
    assert cabecalho == CABECALHO


def test_planilha_nao_traz_o_rotulo_fraco(planilha):
    """Rotulagem as cegas: ver o palpite da Gemini inflaria o Kappa."""
    aba = planilha["rotulagem"]
    cabecalho = [c.value for c in aba[1]]
    assert "rotulo_fraco" not in cabecalho
    todos = " ".join(str(c.value) for linha in aba.iter_rows() for c in linha)
    for classe in CLASSES:
        assert f"rotulo_fraco={classe}" not in todos


def test_planilha_leva_o_texto_original_com_emoji(planilha):
    aba = planilha["rotulagem"]
    textos = [aba.cell(row=r, column=2).value for r in range(2, 5)]
    assert "amei 😂" in textos
    assert "amei risos" not in textos


def test_coluna_rotulo_chega_vazia(planilha):
    aba = planilha["rotulagem"]
    for r in range(2, 5):
        assert aba.cell(row=r, column=3).value in (None, "")


def test_planilha_tem_validacao_por_lista_nas_tres_classes(planilha):
    aba = planilha["rotulagem"]
    formulas = [dv.formula1 for dv in aba.data_validations.dataValidation]
    assert any(all(classe in f for classe in CLASSES) for f in formulas)


def test_planilha_tem_aba_de_instrucoes(planilha):
    assert "instrucoes" in planilha.sheetnames
    texto = " ".join(
        str(c.value) for linha in planilha["instrucoes"].iter_rows() for c in linha if c.value
    )
    assert "NAO consulte os outros avaliadores" in texto
    assert "id_comentario" in texto


# ------------------------------- pool do sorteio (Secao 9 do manual)


def test_o_sorteio_so_enxerga_quem_esta_sem_particao():
    """A regra da Secao 9 mora no SQL: amostra nova = comentario que ninguem viu.

    Sem o `split IS NULL`, uma segunda rodada (Kappa < 0,60) devolveria parte dos
    mesmos comentarios e o novo Kappa mediria a memoria dos avaliadores, nao o
    manual reescrito.
    """
    sql = inspect.getsource(carregar_sorteaveis)
    assert "e.split IS NULL" in sql
    assert "e.rotulo_fraco IS NOT NULL" in sql


def test_a_populacao_de_cochran_ignora_a_particao():
    """N e o corpus rotulado inteiro: ele nao encolhe porque uma rodada ja gastou parte."""
    sql = inspect.getsource(contar_rotulados)
    assert "e.rotulo_fraco IS NOT NULL" in sql
    assert "split" not in sql.split('"""')[2]


def test_refazer_e_nova_rodada_nao_convivem():
    """Um descarta a amostra anterior, o outro a preserva. Juntos, nao querem dizer nada."""
    with pytest.raises(SystemExit, match="opostos"):
        asyncio.run(sortear(id_execucao=1, refazer=True, nova_rodada=True))
